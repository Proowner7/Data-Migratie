#!/usr/bin/env python3
"""Synchroniseer nieuwe debiteuren uit een exportbestand naar een werkbestand."""

from __future__ import annotations

import argparse
import csv
import json
import logging
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, Iterable, List, Sequence, Tuple

from openpyxl import Workbook, load_workbook


@dataclass
class SyncResult:
    timestamp: str
    source_total: int
    existing_total: int
    inserted: int
    skipped_existing: int
    duplicate_source: int
    conflicts: int
    empty_key_rows: int
    output_workbook: str
    report_file: str
    conflict_file: str


def _clean(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _norm(value: object) -> str:
    return _clean(value).upper()


def _parse_csv(path: Path) -> Tuple[List[str], List[Dict[str, str]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        headers = [h.strip() for h in (reader.fieldnames or [])]
        rows = []
        for row in reader:
            rows.append({k.strip(): _clean(v) for k, v in row.items() if k is not None})
    return headers, rows


def _parse_xlsx(path: Path, sheet_name: str | None = None) -> Tuple[List[str], List[Dict[str, str]]]:
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        raise ValueError(f"Bestand '{path}' bevat geen header.")
    headers = [_clean(v) for v in header_row]
    rows: List[Dict[str, str]] = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = {headers[i]: _clean(values[i] if i < len(values) else "") for i in range(len(headers))}
        if any(_clean(v) for v in row.values()):
            rows.append(row)
    return headers, rows


def read_table(path: Path, sheet_name: str | None = None) -> Tuple[List[str], List[Dict[str, str]]]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return _parse_xlsx(path, sheet_name=sheet_name)
    if suffix == ".csv":
        return _parse_csv(path)
    raise ValueError(f"Niet-ondersteund bestandstype voor '{path}'. Gebruik .xlsx/.xlsm/.csv.")


def validate_columns(headers: Sequence[str], required: Iterable[str], path: Path) -> None:
    header_set = set(headers)
    missing = [col for col in required if col not in header_set]
    if missing:
        raise ValueError(f"Ontbrekende kolommen in '{path}': {', '.join(missing)}")


def resolve_key(row: Dict[str, str], key_columns: Sequence[str]) -> Tuple[str, str]:
    """Geef de eerste niet-lege sleutel terug als (kolomnaam, genormaliseerde waarde)."""
    for col in key_columns:
        candidate = _norm(row.get(col, ""))
        if candidate:
            return col, candidate
    return "", ""


def has_conflict(existing: Dict[str, str], incoming: Dict[str, str], shared_columns: Sequence[str]) -> bool:
    """True als beide records in een gedeelde kolom verschillende niet-lege waarden hebben."""
    for col in shared_columns:
        old = _norm(existing.get(col, ""))
        new = _norm(incoming.get(col, ""))
        if old and new and old != new:
            return True
    return False


def write_workbook(path: Path, headers: Sequence[str], rows: Sequence[Dict[str, str]], sheet_name: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(list(headers))
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    wb.save(path)


def write_csv(path: Path, headers: Sequence[str], rows: Sequence[Dict[str, str]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(headers))
        writer.writeheader()
        for row in rows:
            writer.writerow({h: row.get(h, "") for h in headers})


def sync_debtors(
    source_path: Path,
    work_path: Path,
    output_dir: Path,
    key_columns: Sequence[str],
    required_work_columns: Sequence[str],
    source_sheet: str | None = None,
    work_sheet: str | None = None,
) -> SyncResult:
    src_headers, src_rows = read_table(source_path, sheet_name=source_sheet)
    work_headers, work_rows = read_table(work_path, sheet_name=work_sheet)

    common_keys = [col for col in key_columns if col in src_headers and col in work_headers]
    if not common_keys:
        raise ValueError(
            "Geen gedeelde sleutelkolommen gevonden tussen bron en werkbestand. "
            f"Gevraagd: {', '.join(key_columns)}"
        )

    validate_columns(work_headers, required_work_columns, work_path)

    output_headers = list(work_headers)
    for h in src_headers:
        if h not in output_headers:
            output_headers.append(h)

    existing_by_key: Dict[Tuple[str, str], Dict[str, str]] = {}
    output_rows = [dict(row) for row in work_rows]
    for row in output_rows:
        key_col, key_val = resolve_key(row, common_keys)
        if key_val:
            existing_by_key[(key_col, key_val)] = row

    source_seen: set[Tuple[str, str]] = set()
    conflicts_rows: List[Dict[str, str]] = []
    inserted = skipped_existing = duplicate_source = conflicts = empty_key_rows = 0

    for src_row in src_rows:
        key_col, key_val = resolve_key(src_row, common_keys)
        if not key_val:
            empty_key_rows += 1
            conflicts_rows.append(
                {
                    "type": "empty_key",
                    "key_column": "",
                    "key_value": "",
                    "details": "Geen sleutelwaarde gevonden in bronrecord.",
                }
            )
            continue

        record_key = (key_col, key_val)
        if record_key in source_seen:
            duplicate_source += 1
            conflicts_rows.append(
                {
                    "type": "duplicate_source",
                    "key_column": key_col,
                    "key_value": key_val,
                    "details": "Dubbel sleutelrecord in bronbestand.",
                }
            )
            continue
        source_seen.add(record_key)

        current = existing_by_key.get(record_key)
        if current:
            skipped_existing += 1
            shared = [h for h in src_headers if h in work_headers]
            if has_conflict(current, src_row, shared):
                conflicts += 1
                conflicts_rows.append(
                    {
                        "type": "existing_conflict",
                        "key_column": key_col,
                        "key_value": key_val,
                        "details": "Record bestaat al maar bevat afwijkende waarden.",
                    }
                )
            continue

        new_row = {h: _clean(src_row.get(h, "")) for h in output_headers}
        output_rows.append(new_row)
        existing_by_key[record_key] = new_row
        inserted += 1

    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
    run_dir = output_dir / f"run-{timestamp}"
    run_dir.mkdir(parents=True, exist_ok=True)

    output_workbook = run_dir / "debiteuren-updated.xlsx"
    write_workbook(output_workbook, output_headers, output_rows, sheet_name=work_sheet or "Debiteuren")

    conflict_file = run_dir / "conflicts.csv"
    write_csv(conflict_file, ["type", "key_column", "key_value", "details"], conflicts_rows)

    report_file = run_dir / "report.json"
    result = SyncResult(
        timestamp=timestamp,
        source_total=len(src_rows),
        existing_total=len(work_rows),
        inserted=inserted,
        skipped_existing=skipped_existing,
        duplicate_source=duplicate_source,
        conflicts=conflicts,
        empty_key_rows=empty_key_rows,
        output_workbook=str(output_workbook),
        report_file=str(report_file),
        conflict_file=str(conflict_file),
    )
    report_file.write_text(json.dumps(result.__dict__, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Voeg automatisch nieuwe debiteuren uit een Exact-export toe aan een werkbestand."
    )
    parser.add_argument("--source", required=True, help="Pad naar Exact exportbestand (.xlsx/.csv).")
    parser.add_argument("--workbook", required=True, help="Pad naar bestaand werkbestand (.xlsx/.csv).")
    parser.add_argument("--output-dir", required=True, help="Map voor output per run.")
    parser.add_argument(
        "--key-columns",
        default="Debiteurnummer,KvK,Relatiecode",
        help="Komma-gescheiden sleutelkolommen in prioriteitsvolgorde.",
    )
    parser.add_argument(
        "--required-work-columns",
        default="",
        help="Komma-gescheiden verplichte kolommen in het werkbestand.",
    )
    parser.add_argument("--source-sheet", default=None, help="Optionele sheetnaam voor bronbestand.")
    parser.add_argument("--work-sheet", default=None, help="Optionele sheetnaam voor werkbestand.")
    return parser.parse_args()


def main() -> int:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    args = parse_args()
    source = Path(args.source).expanduser().resolve()
    workbook = Path(args.workbook).expanduser().resolve()
    output_dir = Path(args.output_dir).expanduser().resolve()
    key_columns = [c.strip() for c in args.key_columns.split(",") if c.strip()]
    required_work_columns = [c.strip() for c in args.required_work_columns.split(",") if c.strip()]

    try:
        result = sync_debtors(
            source_path=source,
            work_path=workbook,
            output_dir=output_dir,
            key_columns=key_columns,
            required_work_columns=required_work_columns,
            source_sheet=args.source_sheet,
            work_sheet=args.work_sheet,
        )
    except (ValueError, FileNotFoundError, PermissionError):  # pragma: no cover - CLI feedback
        logging.exception("Synchronisatie mislukt.")
        return 1

    logging.info("Run gereed. Nieuwe debiteuren toegevoegd: %s", result.inserted)
    logging.info("Output workbook: %s", result.output_workbook)
    logging.info("Rapport: %s", result.report_file)
    logging.info("Conflicts: %s", result.conflict_file)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
