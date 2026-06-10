import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from debtor_sync import sync_debtors


def _write_workbook(path: Path, headers: list[str], rows: list[list[str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Debiteuren"
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(path)


def _read_rows(path: Path) -> list[list[str]]:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    return [list(r) for r in ws.iter_rows(values_only=True)]


def test_adds_only_new_debtors(tmp_path: Path) -> None:
    work = tmp_path / "work.xlsx"
    src = tmp_path / "src.xlsx"
    _write_workbook(
        work,
        ["Debiteurnummer", "Naam", "KvK"],
        [["100", "Alpha BV", "111"], ["101", "Beta BV", "222"]],
    )
    _write_workbook(
        src,
        ["Debiteurnummer", "Naam", "KvK"],
        [["101", "Beta BV", "222"], ["102", "Gamma BV", "333"]],
    )

    result = sync_debtors(
        source_path=src,
        work_path=work,
        output_dir=tmp_path / "out",
        key_columns=["Debiteurnummer", "KvK", "Relatiecode"],
        required_work_columns=["Debiteurnummer", "Naam"],
    )

    assert result.inserted == 1
    rows = _read_rows(Path(result.output_workbook))
    assert len(rows) == 4
    assert rows[-1][0] == "102"


def test_reports_duplicate_and_conflict(tmp_path: Path) -> None:
    work = tmp_path / "work.xlsx"
    src = tmp_path / "src.xlsx"
    _write_workbook(work, ["Debiteurnummer", "Naam"], [["200", "Omega BV"]])
    _write_workbook(
        src,
        ["Debiteurnummer", "Naam"],
        [["200", "Omega Nieuw"], ["201", "Delta BV"], ["201", "Delta BV duplicate"]],
    )

    result = sync_debtors(
        source_path=src,
        work_path=work,
        output_dir=tmp_path / "out",
        key_columns=["Debiteurnummer"],
        required_work_columns=["Debiteurnummer"],
    )

    assert result.inserted == 1
    assert result.conflicts == 1
    assert result.duplicate_source == 1


def test_counts_empty_key_rows(tmp_path: Path) -> None:
    work = tmp_path / "work.xlsx"
    src = tmp_path / "src.xlsx"
    _write_workbook(work, ["Debiteurnummer", "Naam"], [["300", "Exist BV"]])
    _write_workbook(src, ["Debiteurnummer", "Naam"], [["", "No Key BV"]])

    result = sync_debtors(
        source_path=src,
        work_path=work,
        output_dir=tmp_path / "out",
        key_columns=["Debiteurnummer"],
        required_work_columns=["Debiteurnummer"],
    )

    assert result.inserted == 0
    assert result.empty_key_rows == 1
